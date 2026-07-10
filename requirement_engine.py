import os
import re

import pandas as pd
import matplotlib.pyplot as plt

from docx import Document

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import getSampleStyleSheet

from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity


# ==========================================================
# CONFIGURATION
# ==========================================================

AMBIGUOUS_WORDS = {
    "fast",
    "quick",
    "easy",
    "efficient",
    "user-friendly",
    "good",
    "best",
    "appropriate",
    "robust"
}

PERFORMANCE_PATTERNS = [
    r"\d+\s*ms",
    r"\d+\s*milliseconds",
    r"\d+\s*sec",
    r"\d+\s*seconds"
]

AUTH_KEYWORDS = [
    "login",
    "authentication",
    "authorization",
    "oauth",
    "sso"
]

ENCRYPTION_KEYWORDS = [
    "encryption",
    "encrypted",
    "tls",
    "ssl",
    "https"
]


# ==========================================================
# LOAD EMBEDDING MODEL
# ==========================================================

try:
    model = SentenceTransformer("all-MiniLM-L6-v2")
except Exception:
    model = None


# ==========================================================
# AMBIGUITY CHECK
# ==========================================================

def check_ambiguity(requirements):

    issues = []

    for req in requirements:

        found = []

        for word in AMBIGUOUS_WORDS:

            if word.lower() in req.lower():
                found.append(word)

        if found:

            issues.append(
                f"Ambiguous wording detected: "
                f"'{req}' ({', '.join(found)})"
            )

    return issues


# ==========================================================
# TESTABILITY CHECK
# ==========================================================

def check_testability(requirements):

    issues = []

    for req in requirements:

        if not re.search(r"\d+", req):

            issues.append(
                f"Requirement lacks measurable criteria: "
                f"'{req}'"
            )

    return issues


# ==========================================================
# DUPLICATE CHECK
# ==========================================================

def check_duplicates(requirements):

    issues = []

    seen = {}

    for req in requirements:

        key = req.lower().strip()

        if key in seen:

            issues.append(
                f"Duplicate requirement detected: '{req}'"
            )

        seen[key] = True

    return issues


# ==========================================================
# SEMANTIC DUPLICATES
# ==========================================================

def check_semantic_duplicates(requirements):

    if not model or len(requirements) < 2:
        return []

    issues = []

    embeddings = model.encode(requirements)

    similarity_matrix = cosine_similarity(embeddings)

    for i in range(len(requirements)):
        for j in range(i + 1, len(requirements)):

            if similarity_matrix[i][j] > 0.85:

                issues.append(
                    f"Possible semantic duplicate:\n"
                    f"'{requirements[i]}'\n"
                    f"'{requirements[j]}'"
                )

    return issues


# ==========================================================
# COMPLETENESS
# ==========================================================

def check_completeness(requirements):

    issues = []

    if len(requirements) < 3:

        issues.append(
            "Requirement set appears incomplete."
        )

    return issues


# ==========================================================
# PERFORMANCE EXTRACTION
# ==========================================================

def extract_response_time(req):

    match = re.search(
        r"(\d+)\s*(ms|milliseconds|sec|seconds)",
        req.lower()
    )

    if not match:
        return None

    value = int(match.group(1))
    unit = match.group(2)

    if "ms" in unit:
        return value / 1000

    return value


def extract_volume(req):

    match = re.search(
        r"([\d,]+)\s*(records|transactions|users)",
        req.lower()
    )

    if not match:
        return None

    return int(
        match.group(1).replace(",", "")
    )


# ==========================================================
# CONFLICTS
# ==========================================================

def check_conflicts(requirements):

    issues = []

    response_time = None
    volume = None

    for req in requirements:

        rt = extract_response_time(req)

        if rt:
            response_time = rt

        vol = extract_volume(req)

        if vol:
            volume = vol

    if response_time and volume:

        if response_time < 2 and volume > 500000:

            issues.append(
                f"Potential conflict detected: "
                f"{volume:,} records within "
                f"{response_time} second(s)"
            )

    return issues


# ==========================================================
# SECURITY CHECKS
# ==========================================================

def check_authentication(requirements):

    content = " ".join(requirements).lower()

    if any(
        keyword in content
        for keyword in AUTH_KEYWORDS
    ):
        return []

    return [
        "No authentication requirement detected."
    ]


def check_encryption(requirements):

    content = " ".join(requirements).lower()

    if any(
        keyword in content
        for keyword in ENCRYPTION_KEYWORDS
    ):
        return []

    return [
        "No encryption requirement detected."
    ]


# ==========================================================
# PERFORMANCE CHECK
# ==========================================================

def check_performance(requirements):

    for req in requirements:

        for pattern in PERFORMANCE_PATTERNS:

            if re.search(pattern, req.lower()):

                return []

    return [
        "No measurable performance target specified."
    ]


# ==========================================================
# TESTABILITY SCORE
# ==========================================================

def calculate_testability(req):

    score = 0

    lower = req.lower()

    if "shall" in lower or "must" in lower:
        score += 20

    if re.search(r"\d+", req):
        score += 30

    if any(word in lower for word in AMBIGUOUS_WORDS):
        score -= 20
    else:
        score += 20

    if re.search(r"\d+\s*(sec|seconds|ms|%)", lower):
        score += 30

    return max(0, min(score, 100))


# ==========================================================
# RTM GENERATOR
# ==========================================================

def generate_rtm(requirements):

    rows = []

    for idx, req in enumerate(
        requirements,
        start=1
    ):

        rows.append(
            {
                "Requirement ID": f"REQ-{idx:03}",
                "Requirement": req,
                "Module": "TBD",
                "Test Case": f"TC-{idx:03}",
                "Status": "Open"
            }
        )

    return pd.DataFrame(rows)


# ==========================================================
# DASHBOARD
# ==========================================================

def create_dashboard(result):

    os.makedirs("static", exist_ok=True)

    categories = []
    values = []

    for category, issues in result["issues"].items():

        categories.append(category)
        values.append(len(issues))

    plt.figure(figsize=(10, 6))

    plt.bar(
        categories,
        values,
        color="steelblue"
    )

    plt.title("Requirements Quality Dashboard")
    plt.ylabel("Issue Count")
    plt.xticks(rotation=25)

    chart_file = "static/dashboard.png"

    plt.tight_layout()
    plt.savefig(chart_file)
    plt.close()

    return chart_file


# ==========================================================
# WORD EXPORT
# ==========================================================

def export_word(result, filename):

    doc = Document()

    doc.add_heading(
        "Requirements Analysis Report",
        0
    )

    doc.add_paragraph(
        f"Satisfaction Score: "
        f"{result['satisfaction']}"
    )

    doc.add_heading(
        "Issues",
        level=1
    )

    for category, items in result["issues"].items():

        doc.add_heading(
            category.capitalize(),
            level=2
        )

        for item in items:
            doc.add_paragraph(item)

    doc.save(filename)


# ==========================================================
# PDF EXPORT
# ==========================================================

def export_pdf(result, filename):

    doc = SimpleDocTemplate(filename)

    styles = getSampleStyleSheet()

    story = []

    story.append(
        Paragraph(
            "Requirements Analysis Report",
            styles["Title"]
        )
    )

    story.append(Spacer(1, 10))

    story.append(
        Paragraph(
            f"Satisfaction Score: "
            f"{result['satisfaction']}",
            styles["Normal"]
        )
    )

    story.append(Spacer(1, 10))

    for category, items in result["issues"].items():

        story.append(
            Paragraph(
                category.capitalize(),
                styles["Heading2"]
            )
        )

        for item in items:

            story.append(
                Paragraph(
                    item,
                    styles["BodyText"]
                )
            )

    doc.build(story)


# ==========================================================
# MAIN ANALYSIS
# ==========================================================

def analyze_requirements(requirements):

    ambiguity = check_ambiguity(requirements)

    testability_issues = check_testability(
        requirements
    )

    duplicates = check_duplicates(
        requirements
    )

    semantic_duplicates = (
        check_semantic_duplicates(
            requirements
        )
    )

    completeness = check_completeness(
        requirements
    )

    conflicts = check_conflicts(
        requirements
    )

    authentication = (
        check_authentication(
            requirements
        )
    )

    encryption = check_encryption(
        requirements
    )

    performance = check_performance(
        requirements
    )

    testability_scores = []

    for req in requirements:

        testability_scores.append(
            {
                "requirement": req,
                "score":
                calculate_testability(req)
            }
        )

    rtm = generate_rtm(requirements)

    issues = {
        "ambiguity": ambiguity,
        "testability": testability_issues,
        "duplicates": duplicates,
        "semantic_duplicates": semantic_duplicates,
        "completeness": completeness,
        "conflicts": conflicts,
        "authentication": authentication,
        "encryption": encryption,
        "performance": performance
    }

    total_issues = sum(
        len(v)
        for v in issues.values()
    )

    satisfaction = max(
        0,
        round(1 - (total_issues * 0.05), 2)
    )

    result = {
        "issues": issues,
        "satisfaction": satisfaction,
        "testability_scores": testability_scores,
        "rtm": rtm.to_dict("records")
    }

    chart = create_dashboard(result)

    result["dashboard"] = chart

    return result